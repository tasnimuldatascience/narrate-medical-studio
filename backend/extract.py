"""Inspect OOXML without executing macros or resolving external resources."""
import re
import zipfile
from pathlib import PurePosixPath
from lxml import etree
from pptx import Presentation

MAX_EXPANDED = 512 * 1024 * 1024

def validate_svg(data):
    """Accept static, self-contained SVGs; reject active/external features."""
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    root = etree.fromstring(data, parser)
    tree = root.getroottree()
    error = 'This presentation contains an SVG with active or external content. Replace that image with a static PNG before importing.'
    if tree.docinfo.doctype or tree.xpath('//processing-instruction()'):
        raise ValueError(error)

    def check_css(value):
        # Reject escapes and at-rules rather than interpreting obfuscated URLs.
        value = re.sub(r'/\*.*?\*/', '', value, flags=re.S)
        if '\\' in value or '@' in value or re.search(r'expression\s*\(',value,re.I):
            raise ValueError(error)
        value = re.sub(r'''url\s*\(\s*(["']?)(#[\w:.-]+)\1\s*\)''','',value,flags=re.I)
        if re.search(r'url\s*\(',value,re.I):
            raise ValueError(error)

    for node in root.iter():
        if not isinstance(node.tag,str):
            continue
        name = etree.QName(node).localname.lower()
        if etree.QName(node).namespace not in (None,'http://www.w3.org/2000/svg'):
            raise ValueError(error)
        if name in {'script','foreignobject','animate','animatemotion','animatetransform','set','handler'}:
            raise ValueError(error)
        if name == 'style':
            check_css(''.join(node.itertext()))
        for attribute,value in node.attrib.items():
            key = etree.QName(attribute).localname.lower()
            if key.startswith('on') or key == 'base':
                raise ValueError(error)
            if key in {'href','src'} and value.strip() and not value.strip().startswith('#'):
                raise ValueError(error)
            # Paint server URLs also occur in fill/stroke/filter attributes.
            check_css(value)


def validate(path):
    if not zipfile.is_zipfile(path):
        raise ValueError('This is not a valid PowerPoint file.')
    with zipfile.ZipFile(path) as z:
        entries = z.infolist()
        if len(entries) > 15000 or sum(i.file_size for i in entries) > MAX_EXPANDED:
            raise ValueError('The presentation exceeds the safe extraction limit (512 MB).')
        if len({i.filename for i in entries}) != len(entries):
            raise ValueError('The presentation contains duplicate archive entries.')
        for i in entries:
            p = PurePosixPath(i.filename.replace('\\', '/'))
            if p.is_absolute() or '..' in p.parts or ':' in i.filename:
                raise ValueError('The presentation contains an unsafe archive path.')
            if i.flag_bits & 1:
                raise ValueError('Password-protected presentations are not supported.')
        if 'ppt/presentation.xml' not in z.namelist():
            raise ValueError('The archive does not contain a PowerPoint presentation.')

def sanitize(source, target):
    """Remove executable/embedded content and all external relationships before rendering."""
    validate(source)
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as dst:
        content_types = etree.fromstring(src.read('[Content_Types].xml'),parser) if '[Content_Types].xml' in src.namelist() else []
        if isinstance(content_types,etree._Element) and content_types.getroottree().docinfo.doctype:
            raise ValueError('XML document type declarations are not supported.')
        svg_extensions = {'svg','svgz'}
        svg_parts = set()
        for item in content_types:
            if item.get('ContentType','').lower() == 'image/svg+xml':
                if item.get('Extension'):
                    svg_extensions.add(item.get('Extension').lower())
                if item.get('PartName'):
                    svg_parts.add(item.get('PartName').lstrip('/'))
        removed = {}
        for name in src.namelist():
            if not name.endswith('.rels'):
                continue
            path = PurePosixPath(name)
            owner = str(path.parent.parent / path.name[:-5]) if name != '_rels/.rels' else ''
            relationships = etree.fromstring(src.read(name), parser)
            if relationships.getroottree().docinfo.doctype:
                raise ValueError('XML document type declarations are not supported.')
            removed[owner] = {
                rel.get('Id') for rel in relationships
                if rel.get('TargetMode', '').strip().lower() == 'external'
                or any(v in rel.get('Type', '').lower().rsplit('/', 1)[-1] for v in ('vba', 'oleobject', 'control', 'package'))
            }
        for entry in src.infolist():
            name = entry.filename
            if any(v in name.lower() for v in ('vbaproject', 'vbadata', 'activex/', 'embeddings/')):
                continue
            data = src.read(name)
            if name.rsplit('.',1)[-1].lower() in svg_extensions or name in svg_parts:
                validate_svg(data)
            if name.lower().endswith(('.xml', '.rels')):
                # OOXML does not need DTDs. Reject entity declarations rather
                # than passing them to PowerPoint or downstream XML parsers.
                document = etree.fromstring(data, parser)
                if document.getroottree().docinfo.doctype:
                    raise ValueError('XML document type declarations are not supported.')
            if name.endswith('.rels'):
                root = etree.fromstring(data, parser)
                for rel in list(root):
                    kind = rel.get('Type', '').lower().rsplit('/', 1)[-1]
                    if rel.get('TargetMode', '').strip().lower() == 'external' or any(v in kind for v in ('vba', 'oleobject', 'control', 'package')):
                        root.remove(rel)
                data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            elif name == '[Content_Types].xml':
                root = etree.fromstring(data, parser)
                for item in list(root):
                    ct = item.get('ContentType', '')
                    if any(v in ct.lower() for v in ('vbaproject', 'vbadata', 'activex', 'oleobject')):
                        root.remove(item)
                    elif ct == 'application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml':
                        item.set('ContentType', 'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml')
                data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            elif name in removed and removed[name]:
                root = etree.fromstring(data, parser)
                for node in list(root.iter()):
                    for key, value in list(node.attrib.items()):
                        if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in removed[name]:
                            # A dangling hyperlink ID makes otherwise valid decks
                            # unreadable in PowerPoint. Remove its action node.
                            if etree.QName(node).localname in ('hlinkClick', 'hlinkHover', 'externalData', 'oleObj', 'control'):
                                parent = node.getparent()
                                if parent is not None:
                                    parent.remove(node)
                                break
                            del node.attrib[key]
                data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            dst.writestr(name, data)

def shapes(items):
    for shape in items:
        if shape.shape_type == 6:
            yield from shapes(shape.shapes)
        else:
            yield shape

def extract(path):
    validate(path)
    prs = Presentation(path)
    if len(prs.slides) > 500:
        raise ValueError('Please split presentations with more than 500 slides.')
    result = []
    for i, slide in enumerate(prs.slides, 1):
        text, tables, images = [], [], 0
        heading_candidates = []
        for shape in shapes(slide.shapes):
            if shape.has_text_frame and shape.text.strip():
                text.append(shape.text.strip())
                sizes = [run.font.size or paragraph.font.size or 0
                         for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                # Large step numbers and symbols are decorative, not headings.
                # Keep them in extracted source text, but do not let them beat
                # an alphabetic heading when the slide has no title placeholder.
                if any(character.isalpha() for character in shape.text):
                    heading_candidates.append((max(sizes,default=0),shape.text.strip()))
            if shape.has_table:
                tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
            if shape.shape_type == 13:
                images += 1
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ''
        sources = re.findall(r'https?://[^\s\]]+', notes)
        script = re.sub(r'\[Sources\].*?\[/Sources\]', '', notes, flags=re.S).strip()
        content = '\n\n'.join(text)
        flags = []
        if not content and not tables:
            flags.append('Image-only slide. Review the visual; image text has not been transcribed.')
        elif len(content.split()) < 12:
            flags.append('Limited source text. A reviewer may need to supply the explanation.')
        if tables:
            flags.append('Table present. Verify row/column context and all numerical values.')
        if re.search(r'\d', content):
            flags.append('Numerical content. Check doses, units, ranges, and procedural measurements.')
        if images:
            flags.append('Visual content requires review; local text drafting does not interpret images.')
        if '\ufffd' in content:
            flags.append('Source contains an unreadable character. Correct it against the visual.')
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ''
        if not title and heading_candidates:
            title = max(heading_candidates,key=lambda candidate:candidate[0])[1]
        result.append({'number':i, 'title':title.split('\n')[0][:120] if title else f'Slide {i}',
            'text':content, 'notes':notes, 'tables':tables, 'image_count':images, 'sources':sources,
            'narration':script or content, 'draft_origin':'speaker notes' if script else 'slide text',
            'flags':flags, 'approved':False, 'revision':1, 'reviewer':'', 'approved_at':None,
            'image':False, 'duration':None})
    return result, {'width':prs.slide_width, 'height':prs.slide_height, 'slide_count':len(result)}
