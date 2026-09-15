import hashlib
import math
import threading
import zipfile

import pytest
from lxml import etree
from pptx import Presentation

from backend import extract, media, ppt_animation


def safe_deck(tmp_path):
    original = tmp_path / 'original.pptx'
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(original)
    source = tmp_path / 'safe.pptx'
    extract.sanitize(original, source)
    return source


def test_actual_ooxml_effect_presence_and_source_unchanged(tmp_path):
    source = safe_deck(tmp_path)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    result = ppt_animation.inspect(source)
    assert result['slide_count'] == 1
    assert result['timing_slides'] == result['transition_slides'] == 0
    assert not result['slides'][0]['embedded_media']
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    with zipfile.ZipFile(source) as archive:
        content = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(content['ppt/slides/slide1.xml'])
    etree.SubElement(root, '{%s}transition' % ppt_animation.NS['p'])
    timing = etree.SubElement(root, '{%s}timing' % ppt_animation.NS['p'])
    etree.SubElement(timing, '{%s}cTn' % ppt_animation.NS['p'], presetClass='entr', nodeType='interactiveSeq')
    content['ppt/slides/slide1.xml'] = etree.tostring(root)
    changed = tmp_path / 'metadata-fixture.pptx'
    with zipfile.ZipFile(changed, 'w') as archive:
        for name, data in content.items():
            archive.writestr(name, data)
    record = ppt_animation.inspect(changed)
    assert record['timing_slides'] == record['transition_slides'] == 1
    assert record['slides'][0]['effect_count'] == 1
    assert record['slides'][0]['interactive']


def test_click_effects_are_automated_but_groups_and_finite_repeats_preserved():
    plan = ppt_animation.plan_effects([
        {'trigger': 1, 'duration': 1, 'delay': 0},
        {'trigger': 2, 'duration': 2, 'delay': .2},
        {'trigger': 3, 'duration': .5, 'repeat_count': 2, 'auto_reverse': True},
    ], .5)
    assert plan['clicks_automated'] == 1
    assert plan['effects'][0]['trigger'] == 3
    assert plan['effects'][0]['start'] == .35
    assert plan['effects'][1]['start'] == pytest.approx(.55)
    assert plan['effects'][2]['end'] == pytest.approx(4.55)
    assert plan['minimum_duration'] == pytest.approx(5.55)


@pytest.mark.parametrize('effect', [
    {'trigger': 4, 'duration': 1}, {'trigger': 5, 'duration': 1},
    {'trigger': 1, 'duration': math.inf}, {'trigger': 1, 'duration': 1, 'repeat_count': -1},
])
def test_nonflattenable_or_indefinite_effects_are_explicitly_rejected(effect):
    with pytest.raises(ValueError):
        ppt_animation.plan_effects([effect])


@pytest.mark.parametrize('duration', [0, -.5, math.nan, math.inf, 3601])
def test_invalid_target_duration_never_launches_office(tmp_path, monkeypatch, duration):
    monkeypatch.setattr(media, 'run', lambda *args, **kw: pytest.fail('Worker must not launch'))
    with pytest.raises(ValueError):
        ppt_animation.render_slide(tmp_path / 'safe.pptx', 1, tmp_path / 'out.mp4', duration, threading.Event())


def test_cancelled_export_never_launches_office(tmp_path, monkeypatch):
    cancel = threading.Event()
    cancel.set()
    monkeypatch.setattr(media, 'run', lambda *args, **kw: pytest.fail('Worker must not launch'))
    with pytest.raises(media.Cancelled):
        ppt_animation.render_slide(tmp_path / 'safe.pptx', 1, tmp_path / 'out.mp4', 5, cancel)


def test_missing_native_runtime_is_explicit_and_output_is_not_created(tmp_path, monkeypatch):
    source = safe_deck(tmp_path)
    monkeypatch.setattr(ppt_animation, 'available', lambda: False)
    with pytest.raises(ValueError, match='PowerPoint'):
        ppt_animation.render_slide(source, 1, tmp_path / 'out.mp4', 5, threading.Event())
    assert not (tmp_path / 'out.mp4').exists()


def test_existing_export_is_not_overwritten(tmp_path):
    source = safe_deck(tmp_path)
    output = tmp_path / 'out.mp4'
    output.write_bytes(b'preserve existing video')
    with pytest.raises(ValueError, match='already exists'):
        ppt_animation.render_slide(source, 1, output, 5, threading.Event())
    assert output.read_bytes() == b'preserve existing video'


def test_workspace_retries_office_file_handle_release_without_masking_cancellation(tmp_path, monkeypatch):
    actual = ppt_animation.shutil.rmtree
    attempts = []
    def temporarily_locked(path):
        attempts.append(path)
        if len(attempts) == 1:
            raise PermissionError('Office handle is closing')
        return actual(path)
    monkeypatch.setattr(ppt_animation.shutil, 'rmtree', temporarily_locked)
    monkeypatch.setattr(ppt_animation.time, 'sleep', lambda _: None)
    with pytest.raises(media.Cancelled):
        with ppt_animation.native_workspace(tmp_path) as folder:
            (folder / 'native.mp4').write_bytes(b'partial')
            raise media.Cancelled('cancelled')
    assert len(attempts) == 2
    assert not folder.exists()


@pytest.mark.parametrize('duration', [-1, math.nan, math.inf])
def test_invalid_transition_timing_is_rejected(duration):
    with pytest.raises(ValueError, match='transition'):
        ppt_animation.plan_effects([], duration)
