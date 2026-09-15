import io
import json
import threading
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from backend import app as studio, extract, media, store

def presentation():
    prs=Presentation()
    for title,body in [('Welcome','This is a software validation lesson.'),('Learning together','Pause, reflect, and review the explanation.')]:
        s=prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text=title
        s.placeholders[1].text=body
    data=io.BytesIO()
    prs.save(data)
    return data.getvalue()

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'DATA',tmp_path)
    with TestClient(studio.app) as client:
        yield client

def upload(client):
    response=client.post('/api/projects',files={'file':('example.pptx',presentation())})
    assert response.status_code==201,response.text
    return response.json()

def test_import_and_approval_invalidation(client):
    p=upload(client)
    assert len(p['slides'])==2
    root='/api/projects/'+p['id']
    assert client.post(root+'/export',json={'start':1,'end':2}).status_code==400
    approve={'revision':1,'reviewer':'Test reviewer','acknowledged':True}
    assert client.post(root+'/slides/1/approve',json=approve).status_code==400
    with store.edit(p['id']) as stored:
        stored['slides'][0]['image']=True
    assert client.post(root+'/slides/1/approve',json={**approve,'acknowledged':False}).status_code==400
    assert client.post(root+'/slides/1/approve',json=approve).status_code==200
    r=client.put(root+'/slides/1',json={'revision':1,'narration':'Updated explanation.'})
    assert r.status_code==200
    assert not r.json()['slides'][0]['approved']
    assert r.json()['slides'][0]['revision']==2
    assert client.post(root+'/slides/1/approve',json=approve).status_code==409
    assert client.put(root+'/slides/1',json={'revision':1,'narration':'Stale'}).status_code==409

def test_knowledge_check_validation(client):
    p=upload(client);url=f'/api/projects/{p["id"]}/slides/1'
    check={'question':'What next?','options':['','',''],'correct':0,'explanation':'Review the lesson.'}
    assert client.put(url,json={'revision':1,'narration':'Hello','knowledge':check}).status_code==400
    check['options']=['Review','Skip','Ignore']
    r=client.put(url,json={'revision':1,'narration':'Hello','knowledge':check})
    assert r.status_code==200
    assert r.json()['slides'][0]['knowledge']['correct']==0

def test_locked_during_job(client):
    p=upload(client)
    with store.edit(p['id']) as stored:
        stored['job']={'status':'running'}
    r=client.put(f'/api/projects/{p["id"]}/slides/1',json={'revision':1,'narration':'Changed'})
    assert r.status_code==409

def test_origin_and_download_guards(client):
    assert client.get('/api/projects',headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/api/projects',headers={'Host':'evil.example'}).status_code==400
    p=upload(client)
    assert client.get(f'/api/projects/{p["id"]}/exports/missing/source.pptx').status_code==404
    assert client.get('/api/projects/nonexistent').status_code==404
    assert client.post('/api/projects',files={'file':('bad.pptx',b'not a zip')}).status_code==400
    assert client.post('/api/projects',files={'file':('bad.exe',b'not a zip')}).status_code==400

def test_sanitize_macros_and_external_relationships(tmp_path):
    source=tmp_path/'macro.pptm';target=tmp_path/'safe.pptx'
    with zipfile.ZipFile(io.BytesIO(presentation())) as original,zipfile.ZipFile(source,'w') as z:
        for entry in original.infolist():
            data=original.read(entry.filename)
            if entry.filename=='ppt/_rels/presentation.xml.rels':
                data=data.replace(b'</Relationships>',b'<Relationship Id="evil" Type="vbaProject" Target="vbaProject.bin"/><Relationship Id="ext" Type="image" Target="https://evil.example" TargetMode="External"/></Relationships>')
            z.writestr(entry.filename,data)
        z.writestr('ppt/vbaProject.bin',b'never execute')
    extract.sanitize(source,target)
    with zipfile.ZipFile(target) as z:
        assert 'ppt/vbaProject.bin' not in z.namelist()
        assert b'evil.example' not in z.read('ppt/_rels/presentation.xml.rels')
    assert len(extract.extract(target)[0])==2

def test_reject_zip_traversal(tmp_path):
    source=tmp_path/'bad.pptx'
    with zipfile.ZipFile(source,'w') as z:
        z.writestr('../unsafe',b'bad')
        z.writestr('ppt/presentation.xml',b'bad')
    with pytest.raises(ValueError,match='unsafe'):
        extract.validate(source)

def test_caption_text_and_time(tmp_path):
    text='Use a 1.5 mL sample. Check the 18–20 gauge label, then pause and verify the source.'
    chunks=media.caption_chunks(text)
    assert ' '.join(chunks)==text
    assert media.timestamp(59.9996)=='00:01:00,000'
    media.write_captions([{'start':0,'end':3.4,'text':'Hello {\\pos(0,0)}'}],tmp_path)
    assert '{' not in (tmp_path/'captions.ass').read_text(encoding='utf-8')

def test_interactive_package_escapes_source(tmp_path):
    from backend.lesson import package
    for name in ['presentation.mp4','captions.srt','transcript.txt']:
        (tmp_path/name).write_bytes(b'test')
    package(tmp_path,'</script><script>bad()</script>',[{'number':1,'title':'test','start':0,'end':1}],[])
    html=(tmp_path/'index.html').read_text(encoding='utf-8')
    assert '</script><script>bad()' not in html
    assert 'fetch(' not in html
    with zipfile.ZipFile(tmp_path/'interactive-lesson.zip') as z:
        assert 'index.html' in z.namelist()

def test_export_rejects_unapproved_before_speech(tmp_path):
    with pytest.raises(ValueError,match='approval'):
        media.export({'slides':[{'number':1,'approved':False}]},tmp_path,{'start':1,'end':1},threading.Event(),lambda *a:None)


def test_cancel_queued_job_and_retry(client,monkeypatch):
    pending=[]
    monkeypatch.setattr(studio.pool,'submit',lambda operation:pending.append(operation))
    p=upload(client)
    performed=[]
    studio.submit(p['id'],'test',lambda token:performed.append(True))
    assert client.post(f'/api/projects/{p["id"]}/cancel').status_code==200
    pending.pop()()
    assert not performed
    assert store.get(p['id'])['job']['status']=='cancelled'
    assert p['id'] not in studio.cancellations
    studio.submit(p['id'],'test',lambda token:performed.append(True))
    pending.pop()()
    assert performed==[True]
    assert store.get(p['id'])['job']['status']=='completed'


def test_worker_submission_failure_releases_project(client,monkeypatch):
    p=upload(client)
    def fail(operation):
        raise RuntimeError('executor stopped')
    monkeypatch.setattr(studio.pool,'submit',fail)
    with pytest.raises(RuntimeError,match='executor stopped'):
        studio.submit(p['id'],'test',lambda token:None)
    assert store.get(p['id'])['job']['status']=='failed'
    assert p['id'] not in studio.cancellations


def test_missing_export_file_is_404(client):
    p=upload(client)
    with store.edit(p['id']) as saved:
        saved['exports'].append({'id':'recorded'})
    assert client.get(f'/api/projects/{p["id"]}/exports/recorded/presentation.mp4').status_code==404


def test_knowledge_rejects_equivalent_choices_and_blank_text(client):
    p=upload(client)
    check={'question':'Choose an action','options':['Review',' review ','Continue'],
           'correct':0,'explanation':'Review the source.'}
    url=f'/api/projects/{p["id"]}/slides/1'
    assert client.put(url,json={'revision':1,'narration':'Hello','knowledge':check}).status_code==400
    check.update(question=' ',options=['One','Two','Three'])
    assert client.put(url,json={'revision':1,'narration':'Hello','knowledge':check}).status_code==400


def test_cloud_model_alias_rejected_before_source_submission(monkeypatch):
    from backend import narration
    calls=[]
    class Response:
        def __init__(self,payload):self.payload=payload
        def raise_for_status(self):pass
        def json(self):return self.payload
    class Client:
        def __init__(self,**kwargs):assert kwargs['trust_env'] is False
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def get(self,url):return Response({'models':[{'name':'innocent-alias'}]})
        def post(self,url,json):
            calls.append((url,json))
            assert url.endswith('/api/show')
            return Response({'remote_host':'https://example.invalid','remote_model':'remote'})
    monkeypatch.setattr(narration.httpx,'Client',Client)
    with pytest.raises(ValueError,match='Remote-backed'):
        narration.generate({'text':'private source','notes':'','tables':[]},'innocent-alias','Practitioners')
    assert len(calls)==1
    assert 'private source' not in json.dumps(calls)


def test_project_transaction_rolls_back_on_error(client):
    p=upload(client)
    with pytest.raises(RuntimeError):
        with store.edit(p['id']) as saved:
            saved['name']='Must not persist'
            raise RuntimeError('stop')
    assert store.get(p['id'])['name']==p['name']


def test_reject_xml_external_entity_declarations(tmp_path):
    source=tmp_path/'entity.pptx'
    with zipfile.ZipFile(source,'w') as archive:
        archive.writestr('ppt/presentation.xml',
            '<!DOCTYPE presentation [<!ENTITY secret SYSTEM "file:///private">]><presentation/>')
    with pytest.raises(ValueError,match='document type'):
        extract.sanitize(source,tmp_path/'safe.pptx')


def test_join_audio_pads_at_each_video_boundary(tmp_path):
    import wave
    parts=[]
    for i in range(3):
        path=tmp_path/f'{i}.wav'
        with wave.open(str(path),'wb') as audio:
            audio.setparams((1,2,24000,0,'NONE','not compressed'))
            audio.writeframes(b'\x01\x00'*2000)
        parts.append((path,.125))
    output=tmp_path/'joined.wav'
    media.join_audio(parts,output,threading.Event())
    with wave.open(str(output),'rb') as audio:
        assert audio.getnframes()==9000
        data=audio.readframes(9000)
        for i in range(3):
            assert data[i*6000:i*6000+4000]==b'\x01\x00'*2000
            assert data[i*6000+4000:(i+1)*6000]==b'\x00'*2000


def test_video_audit_rejects_drift(tmp_path,monkeypatch):
    monkeypatch.setattr(media.shutil,'which',lambda command:'ffprobe')
    monkeypatch.setattr(media,'run',lambda *args,**kwargs:json.dumps({'streams':[
        {'codec_type':'video','width':1920,'height':1080,'duration':'10'},
        {'codec_type':'audio','duration':'11'}, {'codec_type':'subtitle'}],
        'format':{'duration':'11'}}))
    with pytest.raises(ValueError,match='timing'):
        media.audit_video(tmp_path,10,threading.Event())


def test_designed_slide_title_uses_heading_instead_of_brand_label(tmp_path):
    from pptx.util import Pt
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    for label,size in [('Brand label',12),('Learning objective',36),('Body explanation',20)]:
        box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1))
        paragraph=box.text_frame.paragraphs[0]
        paragraph.text=label
        paragraph.font.size=Pt(size)
    path=tmp_path/'designed.pptx'
    prs.save(path)
    assert extract.extract(path)[0][0]['title']=='Learning objective'


def test_sanitizer_removes_dangling_hyperlinks_preserves_core_properties(tmp_path):
    prs=Presentation()
    prs.core_properties.title='Keep this metadata'
    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text='External source'
    run=slide.placeholders[1].text_frame.paragraphs[0].add_run()
    run.text='Reference label'
    run.hyperlink.address='https://example.invalid/reference'
    source=tmp_path/'linked.pptx'
    safe=tmp_path/'safe.pptx'
    prs.save(source)
    extract.sanitize(source,safe)
    with zipfile.ZipFile(safe) as archive:
        assert b'hlinkClick' not in archive.read('ppt/slides/slide1.xml')
        assert b'example.invalid' not in archive.read('ppt/slides/_rels/slide1.xml.rels')
        assert b'core-properties' in archive.read('_rels/.rels')
    opened=Presentation(safe)
    assert opened.core_properties.title=='Keep this metadata'
    assert 'Reference label' in opened.slides[0].placeholders[1].text


def test_voice_preview_fixed_script_cache_and_audio_download(client,monkeypatch):
    import wave
    calls=[]
    def synthesize(text,output,engine,rate,cancel,voice):
        calls.append((text,engine,rate,voice))
        with wave.open(str(output),'wb') as audio:
            audio.setparams((1,2,24000,0,'NONE','not compressed'))
            audio.writeframes(b'\0\0'*240)
        return [],.01
    monkeypatch.setattr(media,'synthesize',synthesize)
    settings={'engine':'kokoro','voice':'af_heart','rate':1}
    first=client.post('/api/voice-preview',json={**settings,'text':'Do not synthesize this unreviewed source.'})
    assert first.status_code==200,first.text
    assert len(calls)==1
    assert 'Do not synthesize' not in calls[0][0]
    assert 'A good lesson' in calls[0][0]
    assert calls[0][1:]==('kokoro',1,'af_heart')
    repeated=client.post('/api/voice-preview',json=settings)
    assert repeated.json()==first.json()
    assert len(calls)==1
    downloaded=client.get(first.json()['url'])
    assert downloaded.status_code==200
    assert downloaded.headers['content-type']=='audio/wav'
    assert downloaded.content.startswith(b'RIFF')
    other=client.post('/api/voice-preview',json={**settings,'voice':'bf_emma','rate':.8})
    assert other.status_code==200
    assert other.json()['url']!=first.json()['url']
    assert calls[-1][1:]==('kokoro',.8,'bf_emma')


@pytest.mark.parametrize('settings',[
    {'engine':'remote'}, {'voice':'../../secret'}, {'rate':.69},
    {'rate':1.41}, {'rate':'NaN'}, {'rate':'Infinity'},
])
def test_voice_preview_rejects_invalid_settings_before_synthesis(client,monkeypatch,settings):
    def unexpected(*args,**kwargs):
        pytest.fail('Invalid preview settings reached synthesis')
    monkeypatch.setattr(media,'synthesize',unexpected)
    assert client.post('/api/voice-preview',json=settings).status_code==422


def test_voice_preview_missing_model_is_actionable_and_not_published(client,monkeypatch):
    monkeypatch.setattr(studio.voice,'available',lambda:False)
    response=client.post('/api/voice-preview',json={'engine':'kokoro'})
    assert response.status_code==400
    assert 'not installed' in response.json()['detail']
    assert not list((store.DATA/'voice-previews').glob('*.wav'))


def test_voice_preview_failed_synthesis_can_retry_without_serving_partial_audio(client,monkeypatch):
    def interrupted(text,output,*args):
        output.write_bytes(b'partial')
        raise ValueError('Synthesis failed')
    monkeypatch.setattr(media,'synthesize',interrupted)
    failed=client.post('/api/voice-preview',json={})
    assert failed.status_code==400
    partial=next((store.DATA/'voice-previews').glob('*.part.wav'))
    assert client.get('/api/voice-previews/'+partial.name).status_code==404
    assert client.get('/api/voice-previews/'+partial.name.replace('.part','')).status_code==404
    monkeypatch.setattr(media,'synthesize',lambda text,output,*args:output.write_bytes(b'completed'))
    retry=client.post('/api/voice-preview',json={})
    assert retry.status_code==200
    assert client.get(retry.json()['url']).content==b'completed'


@pytest.mark.parametrize('filename',['source.pptx','a'*64+'.mp4','a'*63+'.wav','a'*64+'.wav'])
def test_voice_sample_unknown_files_are_not_downloadable(client,filename):
    assert client.get('/api/voice-previews/'+filename).status_code==404


def test_visual_cues_edit_revoke_approval_and_save_without_change_preserves_it(client):
    p=upload(client)
    root=f'/api/projects/{p["id"]}/slides/1'
    with store.edit(p['id']) as saved:
        saved['slides'][0]['image']=True
    approval={'revision':1,'reviewer':'Clinical reviewer','acknowledged':True}
    assert client.post(root+'/approve',json=approval).status_code==200
    payload={'revision':1,'narration':p['slides'][0]['narration'],
             'visual_cues':[{'label':'Source detail','x':0,'y':1}]}
    updated=client.put(root,json=payload)
    assert updated.status_code==200,updated.text
    slide=updated.json()['slides'][0]
    assert slide['visual_cues']==[{**payload['visual_cues'][0],'width':.24,'height':.24}]
    assert slide['revision']==2
    assert slide['approved'] is False
    assert slide['reviewer']=='' and slide['approved_at'] is None
    assert client.post(root+'/approve',json=approval).status_code==409
    assert client.post(root+'/approve',json={**approval,'revision':2}).status_code==200
    unchanged=client.put(root,json={**payload,'revision':2}).json()['slides'][0]
    assert unchanged['revision']==2 and unchanged['approved'] is True
    moved=client.put(root,json={**payload,'revision':2,
        'visual_cues':[{'label':'Source detail','x':.5,'y':.5}]}).json()['slides'][0]
    assert moved['revision']==3 and moved['approved'] is False
    removed=client.put(root,json={**payload,'revision':3,'visual_cues':[]}).json()['slides'][0]
    assert removed['revision']==4 and removed['visual_cues']==[]


@pytest.mark.parametrize('cue,status',[
    ({'label':'','x':.5,'y':.5},422),
    ({'label':' \t\n','x':.5,'y':.5},400),
    ({'label':'x'*121,'x':.5,'y':.5},422),
    ({'label':'Point','x':-.001,'y':.5},422),
    ({'label':'Point','x':.5,'y':1.001},422),
    ({'label':'Point','x':'NaN','y':.5},422),
    ({'label':'Point','x':.5,'y':'Infinity'},422),
    ({'label':'Point','x':.5,'y':.5,'width':.019},422),
    ({'label':'Point','x':.5,'y':.5,'height':1.001},422),
])
def test_visual_cues_invalid_values_do_not_change_slide(client,cue,status):
    p=upload(client)
    original=p['slides'][0]
    response=client.put(f'/api/projects/{p["id"]}/slides/1',json={
        'revision':1,'narration':original['narration'],'visual_cues':[cue]})
    assert response.status_code==status,response.text
    assert store.get(p['id'])['slides'][0]==original


def test_visual_cues_exact_limit_and_max_count(client):
    p=upload(client)
    url=f'/api/projects/{p["id"]}/slides/1'
    cues=[{'label':'x'*120,'x':i/5,'y':0} for i in range(6)]
    body={'revision':1,'narration':p['slides'][0]['narration'],'visual_cues':cues}
    response=client.put(url,json=body)
    assert response.status_code==200,response.text
    assert response.json()['slides'][0]['visual_cues']==[{**c,'width':.24,'height':.24} for c in cues]
    response=client.put(url,json={**body,'revision':2,'visual_cues':cues+[cues[0]]})
    assert response.status_code==422
    assert store.get(p['id'])['slides'][0]['revision']==2


def test_visual_cue_legacy_defaults_preserve_approval_but_resize_revokes(client):
    p=upload(client)
    cue={'label':'Existing point','x':.5,'y':.5}
    with store.edit(p['id']) as saved:
        saved['slides'][0].update(image=True,visual_cues=[cue])
    url=f'/api/projects/{p["id"]}/slides/1'
    assert client.post(url+'/approve',json={
        'revision':1,'reviewer':'Clinical reviewer','acknowledged':True}).status_code==200
    payload={'revision':1,'narration':p['slides'][0]['narration'],
             'visual_cues':[{**cue,'width':.24,'height':.24}]}
    unchanged=client.put(url,json=payload)
    assert unchanged.status_code==200,unchanged.text
    assert unchanged.json()['slides'][0]['revision']==1
    assert unchanged.json()['slides'][0]['approved'] is True
    resized=client.put(url,json={**payload,'visual_cues':[{**cue,'width':.5,'height':.2}]})
    assert resized.status_code==200,resized.text
    assert resized.json()['slides'][0]['revision']==2
    assert resized.json()['slides'][0]['approved'] is False


@pytest.mark.parametrize('style',[None,'studio','whiteboard'])
def test_export_visual_style_reaches_renderer_and_recorded_manifest(client,monkeypatch,style):
    p=upload(client)
    root=f'/api/projects/{p["id"]}'
    with store.edit(p['id']) as saved:
        saved['slides'][0]['image']=True
    assert client.post(root+'/slides/1/approve',json={
        'revision':1,'reviewer':'Clinical reviewer','acknowledged':True}).status_code==200
    pending=[]
    monkeypatch.setattr(studio.pool,'submit',lambda operation:pending.append(operation))
    observed=[]
    def export(snapshot,folder,settings,cancel,progress):
        observed.append(settings)
        (folder/'quality.json').write_text(json.dumps({'verified':True}))
        (folder/'chapters.json').write_text('[]')
        with zipfile.ZipFile(folder/'interactive-lesson.zip','w') as bundle:
            bundle.writestr('index.html','<html></html>')
        return 3.5
    monkeypatch.setattr(media,'export',export)
    body={'start':1,'end':1}
    if style is not None:
        body['visual_style']=style
    response=client.post(root+'/export',json=body)
    assert response.status_code==200,response.text
    assert response.json()['job']['status']=='queued'
    pending.pop()()
    project=store.get(p['id'])
    assert project['job']['status']=='completed',project['job']
    assert observed[0]['visual_style']==(style or 'whiteboard')
    record=project['exports'][0]
    manifest=client.get(root+f'/exports/{record["id"]}/manifest.json').json()
    assert manifest['settings']['visual_style']==(style or 'whiteboard')
    assert manifest['slides'][0]['reviewer']=='Clinical reviewer'
    assert manifest['source_sha256']==p['source_sha256']


def test_export_invalid_visual_style_never_queues(client,monkeypatch):
    p=upload(client)
    monkeypatch.setattr(studio.pool,'submit',lambda operation:pytest.fail('Invalid style queued'))
    assert client.post(f'/api/projects/{p["id"]}/export',json={
        'start':1,'end':1,'visual_style':'arbitrary'}).status_code==422
