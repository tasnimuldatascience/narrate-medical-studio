import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from backend import extract


@pytest.mark.parametrize('decoration',['?','↻','03','→','•'])
def test_title_ignores_larger_decorative_text(tmp_path,decoration):
    presentation=Presentation()
    slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    for label,size in [('Make retrieval a conversation',34),(decoration,52),('Recall the idea',22)]:
        box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1))
        paragraph=box.text_frame.paragraphs[0]
        paragraph.text=label
        paragraph.font.size=Pt(size)
    source=tmp_path/'source.pptx'
    presentation.save(source)
    slides,_=extract.extract(source)
    assert slides[0]['title']=='Make retrieval a conversation'
    assert decoration in slides[0]['text']


def test_explicit_title_placeholder_is_preserved(tmp_path):
    presentation=Presentation()
    slide=presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text='2026'
    slide.placeholders[1].text='Alphabetic body text'
    source=tmp_path/'source.pptx'
    presentation.save(source)
    assert extract.extract(source)[0][0]['title']=='2026'
