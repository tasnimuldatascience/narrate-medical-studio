"""API contracts: assessments are revision-bound and cannot grant approval."""
import io
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from backend import app as studio, agent_workflow, store, media, qwen_voice


@pytest.fixture
def setup(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'DATA',tmp_path)
    monkeypatch.setattr(studio.pool,'submit',lambda work:work())
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text='Synthetic teaching';slide.placeholders[1].text='Pause and explain.'
    buffer=io.BytesIO();deck.save(buffer)
    with TestClient(studio.app) as client:
        p=client.post('/api/projects',files={'file':('test.pptx',buffer.getvalue())}).json()
        yield client,p,f'/api/projects/{p["id"]}/slides/1'


def result():
    return {'evaluation':{'state':'review_required','clinical_validation':False,'findings':[]},
        'trace':[{'stage':'source_critic','status':'completed'}],'fingerprint':'test'}


def test_evaluation_keeps_script_and_approval_and_binds_source(setup,monkeypatch):
    client,p,url=setup
    def evaluate(slide,text,model,audience,**kwargs):
        assert slide['source_sha256']==p['source_sha256']
        assert kwargs['run_folder'].parent.name=='agent-runs'
        return result()
    monkeypatch.setattr(agent_workflow,'evaluate',evaluate)
    response=client.post(url+'/evaluate',json={'model':'local:model'}).json()
    saved=response['slides'][0]
    assert response['job']['status']=='completed'
    assert saved['narration']==p['slides'][0]['narration']
    assert saved['revision']==1 and saved['approved'] is False
    assert saved['evaluation']['assessed_revision']==1
    assert saved['agent_run']['status']=='completed'
    edited=client.put(url,json={'revision':1,'narration':'New narration.'}).json()['slides'][0]
    assert edited['evaluation'] is None


def test_failed_evaluation_keeps_previous_report_and_reuses_run(setup,monkeypatch):
    client,p,url=setup
    monkeypatch.setattr(agent_workflow,'evaluate',lambda *a,**k:result())
    prior=client.post(url+'/evaluate',json={'model':'local:model'}).json()['slides'][0]['evaluation']
    folders=[]
    def fail(*args,**kwargs):
        folders.append(kwargs['run_folder']);raise ValueError('Critic returned fabricated evidence quote.')
    monkeypatch.setattr(agent_workflow,'evaluate',fail)
    for _ in range(2):
        response=client.post(url+'/evaluate',json={'model':'local:model'}).json()
        assert response['job']['status']=='failed'
        assert response['slides'][0]['evaluation']==prior
        assert 'fabricated' in response['slides'][0]['evaluation_error']
    assert folders[0]==folders[1]


def test_cancelled_agent_draft_never_commits_candidate(setup,monkeypatch):
    client,p,url=setup
    def cancelled(*args,**kwargs):
        kwargs['cancel'].set()
        return {**result(),'candidate':{'narration':'Should not be saved.','review_flags':[]}}
    monkeypatch.setattr(agent_workflow,'run',cancelled)
    response=client.post(url+'/draft',json={'model':'local:model','workflow':'agents'}).json()
    assert response['job']['status']=='cancelled'
    assert response['slides'][0]['revision']==1
    assert response['slides'][0]['narration']==p['slides'][0]['narration']
    assert response['slides'][0]['agent_run']['status']=='cancelled'


def test_agent_draft_commits_assessment_without_approval(setup,monkeypatch):
    client,p,url=setup
    monkeypatch.setattr(agent_workflow,'run',lambda *a,**k:{**result(),
        'candidate':{'narration':'Pause, then explain the idea.','review_flags':['Compare with source.']}})
    saved=client.post(url+'/draft',json={'model':'local:model','workflow':'agents'}).json()['slides'][0]
    assert saved['revision']==2 and saved['approved'] is False
    assert saved['evaluation']['assessed_revision']==2
    assert saved['ai_flags']==['Compare with source.']


def test_quiz_edit_retains_script_assessment_but_revokes_approval(setup,monkeypatch):
    client,p,url=setup
    monkeypatch.setattr(agent_workflow,'evaluate',lambda *a,**k:result())
    client.post(url+'/evaluate',json={'model':'local:model'})
    payload={'revision':1,'narration':p['slides'][0]['narration'],
        'knowledge':{'question':'What next?','options':['Pause','Skip','Leave'],'correct':0,'explanation':'Pause to reflect.'}}
    saved=client.put(url,json=payload).json()['slides'][0]
    assert saved['evaluation']['assessed_revision']==saved['revision']==2
    assert saved['approved'] is False


@pytest.mark.parametrize('settings',[{'engine':'qwen','voice':'af_heart'}, {'engine':'kokoro','voice':'Aiden'},
    {'engine':'qwen','voice':'Aiden','delivery':'arbitrary instruction'}])
def test_voice_settings_reject_mixed_engines(setup,settings):
    client,_,_=setup
    assert client.post('/api/voice-preview',json=settings).status_code==422


def test_qwen_media_dispatches_preset_without_rewriting_text(tmp_path,monkeypatch):
    captured={}
    def voice(text,output,name,rate,cancel,style):
        captured.update(text=text,voice=name,style=style);return [],1
    monkeypatch.setattr(qwen_voice,'synthesize',voice)
    assert media.synthesize('Usually, many.',tmp_path/'a.wav','qwen',1,None,'Aiden',delivery='engaging')==([],1)
    assert captured=={'text':'Usually, many.','voice':'Aiden','style':'engaging'}
