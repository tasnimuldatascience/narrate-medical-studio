import io
import sys
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

import backend
from backend import app as studio, store, narration


@pytest.fixture
def workflow(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'DATA',tmp_path)
    monkeypatch.setattr(studio.pool,'submit',lambda work:work())
    def read(path,cancel):
        if not path.is_file():
            raise FileNotFoundError('Render the slide before OCR.')
        return {'text':'The diagram has three steps.','engine':'synthetic mock'}
    ocr=SimpleNamespace(extract_text=read,available=lambda:True)
    monkeypatch.setattr(backend,'ocr',ocr,raising=False)
    monkeypatch.setitem(sys.modules,'backend.ocr',ocr)
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text='Learning loop'
    slide.placeholders[1].text='Observe, explain, and retrieve.'
    buffer=io.BytesIO()
    prs.save(buffer)
    with TestClient(studio.app) as client:
        response=client.post('/api/projects',files={'file':('synthetic.pptx',buffer.getvalue())})
        assert response.status_code==201
        project=response.json()
        yield client,project,ocr


def url(project):
    return f'/api/projects/{project["id"]}/slides/1'


def rendered(project):
    folder=store.directory(project['id'])/'slides'
    folder.mkdir(exist_ok=True)
    Image.new('RGB',(160,90),'white').save(folder/'slide-1.png')
    with store.edit(project['id']) as saved:
        saved['slides'][0]['image']=True


def approve(client,project,revision):
    return client.post(url(project)+'/approve',json={
        'revision':revision,'reviewer':'Synthetic reviewer','acknowledged':True})


def test_first_ocr_invalidates_approval_and_requires_transcription_review(workflow):
    client,project,ocr=workflow
    rendered(project)
    assert approve(client,project,1).status_code==200
    response=client.post(url(project)+'/ocr')
    assert response.status_code==200
    saved=response.json()['slides'][0]
    assert saved['visual_text']=='The diagram has three steps.'
    assert saved['source_reviewed'] is False
    assert saved['approved'] is False and saved['reviewer']==''
    assert saved['revision']==2
    assert approve(client,project,2).status_code==400
    checked=client.put(url(project)+'/source',json={
        'revision':2,'visual_text':saved['visual_text'],'source_reviewed':True})
    assert checked.status_code==200
    assert checked.json()['slides'][0]['approved'] is False
    assert approve(client,project,3).status_code==200


def test_repeated_ocr_preserves_edited_source_and_approval(workflow):
    client,project,ocr=workflow
    rendered(project)
    changed=client.put(url(project)+'/source',json={
        'revision':1,'visual_text':'Reviewer-corrected wording.','source_reviewed':True})
    assert changed.status_code==200
    assert approve(client,project,2).status_code==200
    response=client.post(url(project)+'/ocr')
    saved=response.json()['slides'][0]
    assert saved['ocr']['text']=='The diagram has three steps.'
    assert saved['visual_text']=='Reviewer-corrected wording.'
    assert saved['source_reviewed'] is True
    assert saved['revision']==2 and saved['approved'] is True


def test_source_edits_stale_revisions_and_active_job_guard(workflow):
    client,project,ocr=workflow
    rendered(project)
    assert approve(client,project,1).status_code==200
    payload={'revision':1,'visual_text':'Corrected source.','source_reviewed':True}
    response=client.put(url(project)+'/source',json=payload)
    saved=response.json()['slides'][0]
    assert saved['approved'] is False
    assert saved['narration_quality'] is None
    assert saved['narration']==project['slides'][0]['narration']
    assert client.put(url(project)+'/source',json=payload).status_code==409
    unchanged=client.put(url(project)+'/source',json={**payload,'revision':2})
    assert unchanged.json()['slides'][0]['revision']==2
    with store.edit(project['id']) as stored:
        stored['job']={'status':'running'}
    assert client.put(url(project)+'/source',json={**payload,'revision':2}).status_code==409
    assert client.post(url(project)+'/ocr').status_code==409


def test_blank_transcription_cannot_be_marked_reviewed(workflow):
    client,project,ocr=workflow
    assert client.put(url(project)+'/source',json={
        'revision':1,'visual_text':'  ','source_reviewed':True}).status_code==400
    assert store.get(project['id'])['slides'][0]['revision']==1


def test_missing_render_or_ocr_runtime_does_not_mutate_source(workflow):
    client,project,ocr=workflow
    response=client.post(url(project)+'/ocr')
    assert response.json()['job']['status']=='failed'
    assert response.json()['slides'][0]['revision']==1
    rendered(project)
    def missing(path,cancel):
        raise RuntimeError('Local OCR runtime is unavailable.')
    ocr.extract_text=missing
    response=client.post(url(project)+'/ocr')
    assert response.json()['job']['status']=='failed'
    saved=response.json()['slides'][0]
    assert not saved.get('visual_text') and saved['approved'] is False


def test_draft_style_and_quality_are_stored_without_approval(workflow,monkeypatch):
    client,project,ocr=workflow
    calls=[]
    def generate(slide,model,audience,style='conversational',cancel=None):
        calls.append((model,style))
        return narration.Draft(narration='Notice the pattern, then explain it.',
            review_flags=['Synthetic diagnostic.'],quality={'word_count':7})
    monkeypatch.setattr(narration,'generate',generate)
    assert client.post(url(project)+'/draft',json={'model':'local','style':'dramatic'}).status_code==422
    assert not calls
    response=client.post(url(project)+'/draft',json={'model':'local','style':'concise'})
    assert response.json()['job']['status']=='completed'
    saved=response.json()['slides'][0]
    assert calls==[('local','concise')]
    assert saved['narration_quality']=={'word_count':7}
    assert saved['approved'] is False and saved['revision']==2
    # Text-only drafting is allowed, but approval still requires a rendered image.
    assert approve(client,project,2).status_code==400


def test_manual_narration_change_invalidates_prior_ai_diagnostics(workflow):
    client,project,ocr=workflow
    with store.edit(project['id']) as saved:
        saved['slides'][0].update(narration_quality={'unsupported_numbers':[]},ai_flags=['Old draft warning.'])
    response=client.put(url(project),json={'revision':1,'narration':'A different explanation.'})
    assert response.status_code==200
    slide=response.json()['slides'][0]
    assert slide.get('narration_quality') is None
    assert 'Old draft warning.' not in slide.get('ai_flags',[])


def test_quiz_only_edit_keeps_narration_diagnostics(workflow):
    client,project,ocr=workflow
    quality={'unsupported_numbers':[],'word_count':5}
    with store.edit(project['id']) as saved:
        saved['slides'][0].update(narration_quality=quality,ai_flags=['Current draft warning.'])
    response=client.put(url(project),json={'revision':1,'narration':project['slides'][0]['narration'],
        'knowledge':{'question':'What follows observation?','options':['Explanation','Skip','Close'],
                     'correct':0,'explanation':'The source lists explanation next.'}})
    slide=response.json()['slides'][0]
    assert slide['narration_quality']==quality
    assert slide['ai_flags']==['Current draft warning.']


def test_unchecking_source_review_revokes_approval_and_blocks_export(workflow):
    client,project,ocr=workflow
    rendered(project)
    response=client.put(url(project)+'/source',json={
        'revision':1,'visual_text':'Reviewed transcription.','source_reviewed':True})
    assert approve(client,project,2).status_code==200
    response=client.put(url(project)+'/source',json={
        'revision':2,'visual_text':'Reviewed transcription.','source_reviewed':False})
    slide=response.json()['slides'][0]
    assert slide['revision']==3 and not slide['approved']
    assert approve(client,project,3).status_code==400
    assert client.post(f'/api/projects/{project["id"]}/export',json={'start':1,'end':1}).status_code==400


@pytest.mark.parametrize('kind',['ocr','draft'])
def test_cancelled_source_work_never_commits_candidate_or_revokes_prior_review(workflow,monkeypatch,kind):
    client,project,ocr=workflow
    rendered(project)
    assert approve(client,project,1).status_code==200
    before=store.get(project['id'])['slides'][0]
    if kind=='ocr':
        def cancelled(path,cancel):
            cancel.set()
            return {'text':'Cancelled candidate'}
        ocr.extract_text=cancelled
        response=client.post(url(project)+'/ocr')
    else:
        def cancelled(*args,**kwargs):
            studio.cancellations[project['id']].set()
            return narration.Draft(narration='Cancelled candidate.',review_flags=[],quality={})
        monkeypatch.setattr(narration,'generate',cancelled)
        response=client.post(url(project)+'/draft',json={'model':'local'})
    assert response.json()['job']['status']=='cancelled'
    assert response.json()['slides'][0]==before


def image_only(project):
    rendered(project)
    with store.edit(project['id']) as saved:
        saved['slides'][0]['text']=''


def test_auto_ocr_unavailable_fails_before_writer_without_source_changes(workflow,monkeypatch):
    client,project,ocr=workflow
    image_only(project)
    before=store.get(project['id'])['slides'][0]
    ocr.available=lambda:False
    def unexpected(*args,**kwargs):
        pytest.fail('Writer must not run without OCR or manually supplied image text.')
    monkeypatch.setattr(narration,'generate',unexpected)
    response=client.post(url(project)+'/draft',json={'model':'local'})
    assert response.json()['job']['status']=='failed'
    assert 'Set up offline OCR' in response.json()['job']['message']
    assert response.json()['slides'][0]==before


def test_auto_ocr_and_draft_are_committed_together_with_one_revision(workflow,monkeypatch):
    client,project,ocr=workflow
    image_only(project)
    def generate(source_slide,*args,**kwargs):
        assert source_slide['visual_text']=='The diagram has three steps.'
        assert source_slide['source_reviewed'] is False
        during=store.get(project['id'])['slides'][0]
        assert not during.get('visual_text') and during['revision']==1
        return narration.Draft(narration='Explain the three steps.',review_flags=['Check image text.'],quality={})
    monkeypatch.setattr(narration,'generate',generate)
    response=client.post(url(project)+'/draft',json={'model':'local'})
    assert response.json()['job']['status']=='completed'
    saved=response.json()['slides'][0]
    assert saved['revision']==2 and saved['approved'] is False
    assert saved['visual_text']=='The diagram has three steps.'
    assert saved['source_reviewed'] is False
    assert saved['narration']=='Explain the three steps.'
    assert approve(client,project,2).status_code==400


@pytest.mark.parametrize('mode',['failure','cancel'])
def test_auto_ocr_candidate_not_persisted_when_writer_fails_or_is_cancelled(workflow,monkeypatch,mode):
    client,project,ocr=workflow
    image_only(project)
    before=store.get(project['id'])['slides'][0]
    def generate(source_slide,*args,cancel=None,**kwargs):
        assert source_slide['visual_text']=='The diagram has three steps.'
        if mode=='failure':
            raise RuntimeError('Writer failed.')
        cancel.set()
        return narration.Draft(narration='Do not store this cancelled draft.',review_flags=[],quality={})
    monkeypatch.setattr(narration,'generate',generate)
    response=client.post(url(project)+'/draft',json={'model':'local'})
    assert response.json()['job']['status']==('failed' if mode=='failure' else 'cancelled')
    assert response.json()['slides'][0]==before


def test_auto_ocr_retains_empty_detection_warning(workflow,monkeypatch):
    client,project,ocr=workflow
    image_only(project)
    ocr.extract_text=lambda *args:{'text':'','warnings':['No readable text was detected. Transcribe the visible source manually.']}
    monkeypatch.setattr(narration,'generate',lambda *args,**kwargs:
        narration.Draft(narration='Source notes need review.',review_flags=[],quality={}))
    response=client.post(url(project)+'/draft',json={'model':'local'})
    saved=response.json()['slides'][0]
    assert 'No readable text' in saved['ocr']['warnings'][0]
    assert saved['source_reviewed'] is False and saved['approved'] is False


def test_existing_reviewed_image_text_skips_automatic_ocr(workflow,monkeypatch):
    client,project,ocr=workflow
    image_only(project)
    client.put(url(project)+'/source',json={
        'revision':1,'visual_text':'Already reviewed text.','source_reviewed':True})
    def unexpected(*args,**kwargs):
        pytest.fail('Automatic OCR must not replace saved image text.')
    ocr.extract_text=unexpected
    def generate(source_slide,*args,**kwargs):
        assert source_slide['visual_text']=='Already reviewed text.'
        assert source_slide['source_reviewed'] is True
        return narration.Draft(narration='Explain the reviewed text.',review_flags=[],quality={})
    monkeypatch.setattr(narration,'generate',generate)
    response=client.post(url(project)+'/draft',json={'model':'local'})
    assert response.json()['job']['status']=='completed'
    assert response.json()['slides'][0]['source_reviewed'] is True
