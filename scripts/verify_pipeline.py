"""Synthetic diagram-to-lesson validation; never approves clinical material.

No arguments: whiteboard output in data/validation/export for the preview server.
Explicit --style: export-whiteboard or export-studio for side-by-side comparison.
"""
import argparse
import json
import sys
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import extract, media, store
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PAGE_W, PAGE_H = 13.333333, 7.5
INK, MUTED = '243F46', '647C80'
TEAL, CORAL, VIOLET = '197F89', 'BE6955', '7675A4'


def text(slide, label, x, y, w, h, size=24, color=INK, bold=False, centered=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = Inches(.04)
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = label
    paragraph.font.name = 'Segoe UI'
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    if centered:
        paragraph.alignment = PP_ALIGN.CENTER
    return box


def shape(slide, kind, x, y, w, h, fill, stroke=None):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = RGBColor.from_string(fill)
    if stroke:
        item.line.color.rgb = RGBColor.from_string(stroke)
        item.line.width = Pt(1.7)
    else:
        item.line.fill.background()
    return item


def arrow(slide, x, y, w=.58, h=.27):
    return shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, '8CA5A5')


def heading(slide, title, description):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    text(slide, 'A BETTER WAY TO LEARN', .65, .36, 12, .3, 13, TEAL, True)
    text(slide, title, .65, .95, 12, .7, 34, INK, True)
    text(slide, description, .65, 1.78, 12, .6, 20, MUTED)


def cue(label, x, y, w, h):
    """Convert an authored source rectangle into normalized coordinates."""
    return {'label':label,'x':(x+w/2)/PAGE_W,'y':(y+h/2)/PAGE_H,
            'width':w/PAGE_W,'height':h/PAGE_H}


def fixture():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    cue_sets = []
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, 'See the learning loop', 'Observation → Explanation → Retrieval')
    for i,(label,detail,color) in enumerate([
            ('Observe','Notice what is there.',TEAL),
            ('Explain','Connect the details.',CORAL),
            ('Retrieve','Put it in your words.',VIOLET)]):
        x = .65+i*4.15
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,x,3.0,3.65,2.5,'F4F8F6',color)
        shape(slide,MSO_SHAPE.OVAL,x+1.40,3.30,.8,.8,color)
        text(slide,str(i+1),x+1.40,3.33,.8,.6,26,'FFFFFF',True,True)
        text(slide,label,x+.15,4.28,3.35,.55,26,color,True,True)
        text(slide,detail,x+.15,4.92,3.35,.38,17,MUTED,False,True)
        if i<2:
            arrow(slide,x+3.72,4.10,.36,.28)
    text(slide,'Start with the source. Return to it when you need to check.',.7,6.25,11.9,.5,21,MUTED,False,True)
    slide.notes_slide.notes_text_frame.text=(
        'Start by noticing what is on the page. Then connect those details into an explanation. '
        'Finally, look away and put the idea into your own words. That is the loop we will practice.')
    cue_sets.append([cue(label,.55+i*4.15,2.9,3.86,2.75)
                     for i,label in enumerate(['Observe the source','Connect the details','Retrieve in your words'])])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide,'Turn observation into explanation','Notice a pattern. Connect its parts. Say what you see.')
    # An abstract pattern, with no anatomy or clinical information.
    for x,y,r in [(1.28,3.1,.45),(2.46,3.1,.45),(1.88,4.16,.45)]:
        shape(slide,MSO_SHAPE.OVAL,x,y,r,r,TEAL)
    shape(slide,MSO_SHAPE.RECTANGLE,1.70,3.28,.79,.07,'A9CBC7')
    shape(slide,MSO_SHAPE.RECTANGLE,2.07,3.55,.07,.68,'A9CBC7')
    text(slide,'Notice a pattern',.75,5.15,3.5,.5,23,TEAL,True,True)
    arrow(slide,4.00,3.96,.62,.3)
    for row,label in enumerate(['What repeats?','What connects?']):
        shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,4.92,3.00+row*1.04,3.45,.77,'FAF0EA',CORAL)
        text(slide,label,5.04,3.17+row*1.04,3.2,.46,21,CORAL,True,True)
    text(slide,'Name the relationship',4.64,5.15,4.0,.5,23,CORAL,True,True)
    arrow(slide,8.65,3.96,.62,.3)
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,9.53,3.00,3.1,1.82,'F1F0F8',VIOLET)
    text(slide,'“Here is how\nthese parts connect.”',9.70,3.39,2.78,1.17,22,VIOLET,True,True)
    text(slide,'Explain in your words',9.02,5.15,4.0,.5,23,VIOLET,True,True)
    slide.notes_slide.notes_text_frame.text=(
        'Look at the small pattern on the left. What repeats, and what connects? '
        'Use those questions to build your explanation. You do not need to copy the sentence; '
        'say what the relationship means to you.')
    cue_sets.append([cue('Notice the source pattern',.65,2.72,3.65,3.05),
                     cue('Ask what connects',4.65,2.72,3.9,3.05),
                     cue('Explain the relationship',9.02,2.72,4.0,3.05)])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide,'Make retrieval a conversation','Pause → Recall → Check')
    for i,(title,detail,color) in enumerate([
            ('Pause','Look away\nfrom the source.',TEAL),
            ('Recall','Say the idea\nin your own words.',CORAL),
            ('Check','Compare, adjust,\nand try again.',VIOLET)]):
        x = .8+i*4.13
        shape(slide,MSO_SHAPE.OVAL,x+.75,2.95,1.8,1.8,'FFFFFF',color)
        if i==0:
            shape(slide,MSO_SHAPE.RECTANGLE,x+1.31,3.47,.19,.75,color)
            shape(slide,MSO_SHAPE.RECTANGLE,x+1.76,3.47,.19,.75,color)
        elif i==1:
            text(slide,'?',x+1.13,3.19,1.0,1.1,50,color,True,True)
        else:
            text(slide,'↻',x+1.01,3.19,1.25,1.1,49,color,True,True)
        text(slide,title,x+.05,4.98,3.2,.5,26,color,True,True)
        text(slide,detail,x,5.56,3.3,.98,20,MUTED,False,True)
        if i<2:
            arrow(slide,x+3.26,3.68,.60,.3)
    slide.notes_slide.notes_text_frame.text=(
        'Now pause and look away from the source. Can you explain the loop without reading it? '
        'When you are ready, check your version against the diagram. Adjust anything you missed, '
        'then give it another try.')
    cue_sets.append([cue(label,.68+i*4.13,2.75,3.60,3.97)
                     for i,label in enumerate(['Pause and look away','Recall the explanation','Check and try again'])])
    return prs,cue_sets


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--style',choices=['whiteboard','studio'],
        help='Explicit styles use export-STYLE folders; omitted uses whiteboard in export.')
    args=parser.parse_args()
    style=args.style or 'whiteboard'
    folder=store.ROOT/'data'/'validation'
    folder.mkdir(parents=True,exist_ok=True)
    prs,cue_sets=fixture()
    prs.save(folder/'source.pptx')
    extract.sanitize(folder/'source.pptx',folder/'safe.pptx')
    slides,metadata=extract.extract(folder/'safe.pptx')
    cancel=threading.Event()
    media.render(folder,cancel)
    for slide,visual_cues in zip(slides,cue_sets):
        slide.update(approved=True,image=True,visual_cues=visual_cues,
            reviewer='Synthetic software validation fixture',approved_at=store.now())
    # Approval is confined to this in-memory synthetic software fixture.
    slides[2]['knowledge']={
        'question':'After recalling the idea in your own words, what comes next?',
        'options':['Compare it with the source','Only reread the title','Skip the explanation'],
        'correct':0,
        'explanation':'The illustrated loop asks you to check your explanation against the source, adjust it, and try again.'}
    output=folder/(f'export-{args.style}' if args.style else 'export')
    output.mkdir(exist_ok=True)
    length=media.export({'name':'A better way to learn','slides':slides},output,
        {'start':1,'end':3,'engine':'kokoro','voice':'af_heart','rate':.96,
         'burn_captions':True,'visual_style':style},cancel,
        lambda percent,message:print(percent,message,flush=True))
    print(json.dumps({'duration':length,'style':style,'output':str(output/'presentation.mp4'),
        'interactive':str(output/'interactive-lesson.zip'),'quality':str(output/'quality.json'),
        'chapters':str(output/'chapters.json')}),flush=True)


if __name__=='__main__':
    main()
