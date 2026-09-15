"""Real, nonclinical PowerPoint animation fixture; supplied decks are untouched."""
import hashlib
import json
import sys
import threading
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import extract, media, ppt_animation
from backend.ppt_video_worker import new_powerpoint_pid


def main():
    import pythoncom
    import win32com.client
    import win32process
    from PIL import Image, ImageChops, ImageStat
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    folder = Path(__file__).resolve().parents[1] / 'data' / 'validation' / 'native-animation'
    folder.mkdir(parents=True, exist_ok=True)
    original = folder / 'synthetic-animation.pptx'
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.8), Inches(.5), Inches(11), Inches(1))
    title.text = 'Native animation software fixture'
    title.text_frame.paragraphs[0].font.size = Pt(32)
    for x, rgb in [(1.5, (225, 70, 80)), (7.5, (40, 140, 175))]:
        box = slide.shapes.add_shape(1, Inches(x), Inches(2.5), Inches(3.3), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*rgb)
        box.line.fill.background()
    deck.save(original)
    pythoncom.CoInitialize()
    app = presentation = None
    owned = False
    before = set(win32process.EnumProcesses())
    try:
        app = win32com.client.DispatchEx('PowerPoint.Application')
        pid = new_powerpoint_pid(before, app)
        if pid in before:
            raise RuntimeError('Close existing PowerPoint before running the synthetic COM fixture.')
        owned = True
        app.AutomationSecurity = 3
        presentation = app.Presentations.Open(str(original), ReadOnly=False, Untitled=False, WithWindow=False)
        source_slide = presentation.Slides(1)
        source_slide.SlideShowTransition.EntryEffect = 1793  # ppEffectFade
        source_slide.SlideShowTransition.Duration = .5
        effect = source_slide.TimeLine.MainSequence.AddEffect(source_slide.Shapes(2), 10, 0, 1)
        effect.Timing.Duration = 1.2
        effect.Timing.TriggerDelayTime = 1.0
        second = source_slide.TimeLine.MainSequence.AddEffect(source_slide.Shapes(3), 2, 0, 3)
        second.Timing.Duration = 1.2
        second.Timing.TriggerDelayTime = .5
        presentation.Save()
    finally:
        if owned:
            if presentation is not None:
                presentation.Close()
            if app is not None:
                app.Quit()
        pythoncom.CoUninitialize()
    safe = folder / 'safe.pptx'
    extract.sanitize(original, safe)
    source_hash = hashlib.sha256(safe.read_bytes()).hexdigest()
    output = folder / ('native-' + source_hash[:12] + '.mp4')
    result = ppt_animation.render_slide(safe, 1, output, 7.25, threading.Event())
    import shutil
    for stamp, name in [(.15, 'early.png'), (6.8, 'late.png')]:
        media.run([shutil.which('ffmpeg'), '-y', '-ss', str(stamp), '-i', str(output), '-frames:v', '1', str(folder / name)], threading.Event())
    early = Image.open(folder / 'early.png').convert('RGB')
    late = Image.open(folder / 'late.png').convert('RGB')
    motion_delta = sum(ImageStat.Stat(ImageChops.difference(early, late)).mean) / 3
    assert motion_delta > 5, f'Expected visible native motion; mean frame difference was {motion_delta}'
    assert hashlib.sha256(safe.read_bytes()).hexdigest() == source_hash
    result.update(source_sha256=source_hash, motion_frame_difference=motion_delta, output=str(output))
    if '--exercise-cancel' in sys.argv:
        cancel = threading.Event()
        timer = threading.Timer(2, cancel.set)
        timer.start()
        cancelled_output = folder / ('cancelled-' + source_hash[:12] + '.mp4')
        started = time.monotonic()
        try:
            ppt_animation.render_slide(safe, 1, cancelled_output, 60, cancel)
            raise AssertionError('Expected cancellation during the longer native export.')
        except media.Cancelled:
            assert not cancelled_output.exists()
            result['cancellation_seconds'] = time.monotonic() - started
        finally:
            timer.cancel()
        # A too-short narration must fail before exporting/truncating effects.
        try:
            ppt_animation.render_slide(safe, 1, folder / ('too-short-' + source_hash[:12] + '.mp4'), 1, threading.Event())
            raise AssertionError('Expected an explicit native-animation duration error.')
        except RuntimeError as error:
            assert 'Native effects need at least' in str(error)
            result['too_short_duration_rejected'] = True
    (folder / 'verification.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps(result, indent=2), flush=True)


if __name__ == '__main__':
    main()
